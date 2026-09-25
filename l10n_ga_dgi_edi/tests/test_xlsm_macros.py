import base64
import io
import os
import zipfile

import xlsxwriter
from openpyxl import load_workbook

from odoo.tests import tagged
from odoo.tools import file_open_temporary_directory

from ..renderers.base import DeclarationBuilder
from ..renderers.xlsm_template import XlsmTemplateRenderer
from ..renderers.xlsx_builder import XlsxDeclarationBuilder
from .common import F16_IRPP, F16_TCS, GaDeclarationCase

VBA_PROJECT = b'\xd0\xcf\x11\xe0 projet VBA factice du gabarit DGI'


def make_xlsm():
    """Gabarit ``.xlsm`` d'essai : projet VBA (factice), formule « fausse » à écraser, 2e feuille."""
    stream = io.BytesIO()
    book = xlsxwriter.Workbook(stream, {'in_memory': True})
    book.set_vba_name('ThisWorkbook')
    sheet = book.add_worksheet('ID')
    sheet.set_vba_name('Feuil1')
    sheet.write_string('A11', 'TCS')
    sheet.write_formula('C11', '=1+1')
    book.add_worksheet('Annexe')
    book.add_vba_project(io.BytesIO(VBA_PROJECT), is_stream=True)
    book.close()
    return stream.getvalue()


def sheet_xml(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


@tagged('post_install', '-at_install')
class TestXlsmMacros(GaDeclarationCase):
    """F10 (règle d'or 10, patron 9) : gabarit .xlsm rempli puis rouvert, projet VBA intact, valeurs seulement."""

    def test_xlsm_keeps_vba_and_writes_values(self):
        renderer = XlsmTemplateRenderer(make_xlsm())
        self.assertIsInstance(renderer, DeclarationBuilder)
        renderer.header({'ID!B2': 'NIF-TEST'})
        renderer.boxes([('C11', 13_058), ('Annexe!B5', '=FORMULE_INJECTEE'), ('ID!D11', None)])
        renderer.table('Annexe', 9, [('F16', 23_195), ('Autre', 1)], headers=('Salarié', 'IRPP'))
        content = renderer.build()

        files = sheet_xml(content)
        self.assertEqual(files['xl/vbaProject.bin'], VBA_PROJECT)
        self.assertIn(b'macroEnabled', files['[Content_Types].xml'])
        for name, data in files.items():
            if name.startswith('xl/worksheets/'):
                self.assertNotIn(b'<f>', data, f'formule dans {name}')

        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book['ID']['C11'].value, 13_058)
        self.assertEqual(book['ID']['B2'].value, 'NIF-TEST')
        self.assertIsNone(book['ID']['D11'].value)
        self.assertEqual(book['Annexe']['B5'].value, '=FORMULE_INJECTEE')
        self.assertEqual(book['Annexe']['B5'].data_type, 's')
        self.assertEqual([c.value for c in book['Annexe'][10]], ['Salarié', 'IRPP'])
        self.assertEqual([c.value for c in book['Annexe'][11]], ['F16', 23_195])
        invalid = XlsmTemplateRenderer(make_xlsm())
        with self.assertRaises(ValueError):
            invalid.boxes([('pas une cellule', 1)])
        invalid.close()

    def test_declaration_fills_template(self):
        self._f16_slip()
        with file_open_temporary_directory(self.env) as directory:
            path = os.path.join(directory, 'ID_TEST.xlsm')
            with open(path, 'wb') as template:
                template.write(make_xlsm())
            decl_type = self._declaration_type('T_XLSM', template_path=path)
            declaration = self._declaration(decl_type=decl_type)
            declaration.action_compute()
            declaration.with_user(self.declarant).action_validate()
        xlsm = declaration.snapshot_attachment_ids.filtered(lambda a: a.name.endswith('.xlsm'))
        self.assertEqual(xlsm.mimetype.lower(), 'application/vnd.ms-excel.sheet.macroenabled.12')
        content = base64.b64decode(xlsm.datas)
        self.assertEqual(sheet_xml(content)['xl/vbaProject.bin'], VBA_PROJECT)
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book['ID']['C11'].value, F16_TCS)
        self.assertEqual(book['ID']['B2'].value, 'NIF-TEST')
        self.assertEqual(book['ID']['B3'].value, 9)

    def test_new_workbook_without_template(self):
        self._f16_slip()
        declaration = self._declaration()
        declaration.action_compute()
        content, extension = declaration._render_xlsx()
        self.assertEqual(extension, 'xlsx')
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book.sheetnames, ['Déclaration', 'Détails'])
        rows = [[cell.value for cell in row] for row in book['Déclaration'].iter_rows()]
        self.assertIn(['B_IRPP', 'IRPP', F16_IRPP], rows)
        self.assertIn(['Société', self.company.name, None], rows)
        details = [[cell.value for cell in row] for row in book['Détails'].iter_rows()]
        self.assertEqual(details[0], ['Case', 'Salarié / tiers', 'Montant'])
        self.assertIn(['B_IRPP', 'F16', F16_IRPP], details)
        for name, data in sheet_xml(content).items():
            if name.startswith('xl/worksheets/'):
                self.assertNotIn(b'<f>', data)

    def test_builder_interface(self):
        base = DeclarationBuilder()
        for call in (
            lambda: base.header({}),
            lambda: base.boxes([]),
            lambda: base.table('x', 0, []),
            base.build,
        ):
            with self.assertRaises(NotImplementedError):
                call()
        builder = XlsxDeclarationBuilder()
        builder.header({'Société': 'X'})
        builder.boxes(['Case'], [('A', 1)])
        builder.table('Détails', 2, [('A', 1)])
        book = load_workbook(io.BytesIO(builder.build()))
        self.assertEqual(book['Détails'].cell(row=4, column=1).value, 'A')
