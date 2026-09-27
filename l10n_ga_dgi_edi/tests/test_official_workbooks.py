import base64
import io
import warnings
import zipfile
from datetime import date

from openpyxl import load_workbook as _load_workbook

from odoo.tests import tagged
from odoo.tools.misc import file_open

from ..models.generators.das import (
    ID19_MARITAL,
    ID19_NATURES,
    ID21_GENDER,
    ID21_MARITAL,
    ID21_NATIONALITY,
    OFFICIAL_TEMPLATES,
)
from ..renderers.xlsm_saisie import SaisieWorkbook, column_index
from .common import GaDeclarationCase


def load_workbook(*args, **kwargs):
    """Lecture de contrôle : openpyxl ignore les listes déroulantes étendues (x14) — sans effet ici."""
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Data Validation extension')
        return _load_workbook(*args, **kwargs)


YEAR = 2026
MONTHS = [(date(YEAR, m, 1), date(YEAR, m, d)) for m, d in ((1, 31), (2, 28), (3, 31))]
UNCHANGED_EXCEPT = {'xl/workbook.xml', 'xl/styles.xml'}


def template(code):
    with file_open(OFFICIAL_TEMPLATES[code][0], 'rb') as stream:
        return stream.read()


def referentiel_values(content):
    book = load_workbook(io.BytesIO(content), read_only=True)
    return {cell for row in book['Referentiel'].iter_rows(values_only=True) for cell in row if isinstance(cell, str)}


@tagged('post_install', '-at_install')
class TestOfficialWorkbooks(GaDeclarationCase):
    """D-87, FIX 05 : classeurs officiels « edi-annexe » de la DGI remplis sans perte (macros, boutons, listes)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.employee = cls._f16_employee(
            'Titulaire',
            l10n_ga_job_code='E01',
            l10n_ga_level_code='N1',
            private_street='BP 100',
            private_city='Libreville',
            l10n_ga_benefit_housing=True,
            country_id=cls.env.ref('base.ga').id,
        )
        for period in MONTHS:
            cls._f16_slip(cls.employee, period)
        wizard = cls.env['l10n_ga.das.wizard'].create({'company_id': cls.company.id, 'year': YEAR})
        cls.das = cls.env['l10n_ga.declaration'].browse(wizard.action_prepare()['res_id'])
        cls.row = cls.das.detail_ids.filtered(lambda d: d.employee_id == cls.employee).payload

    # --- écrivain XML ------------------------------------------------------------------------------

    def test_writer_keeps_every_other_part(self):
        source = template('ID21')
        content = SaisieWorkbook(source).fill({'C9': 'NIF-X'}, 18, [{'B': 'NIF-1', 'M': 1000}]).build()
        with zipfile.ZipFile(io.BytesIO(source)) as before, zipfile.ZipFile(io.BytesIO(content)) as after:
            self.assertEqual(sorted(before.namelist()), sorted(after.namelist()))
            changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
            self.assertIn('xl/vbaProject.bin', after.namelist())
            sheet = next(name for name in changed if name.startswith('xl/worksheets/'))
            self.assertEqual(changed - {sheet}, UNCHANGED_EXCEPT)
            self.assertIn(b'x14:dataValidation', after.read(sheet))  # listes déroulantes étendues conservées
            self.assertIn(b'fullCalcOnLoad="1"', after.read('xl/workbook.xml'))

    def test_writer_values_dates_and_formulas(self):
        workbook = SaisieWorkbook(template('ID21'))
        workbook.fill({}, 18, [{'C': ' Éric ', 'G': 34, 'K': date(2026, 1, 1), 'M': 1234.5}])
        self.assertEqual((workbook.value('C18'), workbook.value('G18')), (' Éric ', '34'))
        self.assertEqual(workbook.value('M18'), '1234.5')
        self.assertEqual(workbook.value('K18'), str((date(2026, 1, 1) - date(1899, 12, 30)).days))
        with self.assertRaisesRegex(ValueError, 'formule'):
            workbook.write('R18', 1)  # total de ligne du classeur
        with self.assertRaisesRegex(ValueError, 'invalide'):
            workbook.write('18B', 1)
        with self.assertRaisesRegex(ValueError, 'booléenne'):
            workbook.write('B18', True)
        workbook.write('B5000', 'hors des lignes existantes')  # ligne créée à sa place
        workbook.write('C18', '')
        self.assertIsNone(workbook.value('C18'))
        sheet = load_workbook(io.BytesIO(workbook.build()))['SAISIE']
        self.assertEqual(sheet['K18'].value.date(), date(2026, 1, 1))
        self.assertTrue(sheet['K18'].is_date)
        self.assertEqual(sheet['B5000'].value, 'hors des lignes existantes')
        self.assertEqual(sheet['R18'].value, '=SUM(M18:Q18)')
        self.assertEqual(column_index('AG'), 33)
        with self.assertRaisesRegex(ValueError, 'absente'):
            SaisieWorkbook(template('ID21'), sheet_name='INEXISTANTE')

    def test_labels_exist_in_templates(self):
        """Libellés écrits = valeurs des listes du classeur (feuille Referentiel) : sinon le XML serait refusé."""
        id19, id21 = referentiel_values(template('ID19')), referentiel_values(template('ID21'))
        self.assertLessEqual(set(ID19_NATURES.values()) | set(ID19_MARITAL.values()) | {'Non précisé'}, id19)
        self.assertLessEqual(
            set(ID21_NATIONALITY.values()) | set(ID21_GENDER.values()) | set(ID21_MARITAL.values()), id21
        )

    # --- DAS : ID19 et ID21 ------------------------------------------------------------------------

    def _books(self, declaration):
        return {
            name.split('-')[-2]: load_workbook(io.BytesIO(content))['SAISIE']
            for name, content in declaration._l10n_ga_official_workbooks()
        }

    def test_id21_filled_from_frozen_details(self):
        self.assertTrue(self.das.has_official_workbooks)
        sheet = self._books(self.das)['ID21']
        self.assertEqual((sheet['C9'].value, sheet['C11'].value, sheet['C13'].value), ('NIF-TEST', YEAR, 'Annuel'))
        rows = [row for row in range(18, 30) if sheet[f'C{row}'].value]
        self.assertEqual(len(rows), len(self.das.detail_ids.filtered('employee_id')))
        line = next(row for row in rows if sheet[f'C{row}'].value == 'Titulaire')
        self.assertEqual(sheet[f'M{line}'].value, self.row['C1'])
        self.assertEqual(sheet[f'N{line}'].value, self.row['C2'])
        self.assertEqual(sheet[f'T{line}'].value, self.row['C8'])
        self.assertEqual(sheet[f'U{line}'].value, self.row['C10'])
        self.assertEqual(sheet[f'F{line}'].value, 'Gabonais')
        self.assertEqual(sheet[f'K{line}'].value.date(), date(YEAR, 1, 1))
        self.assertEqual(sheet[f'R{line}'].value, f'=SUM(M{line}:Q{line})')  # formule du classeur intacte

    def test_id19_rows_by_nature(self):
        sheet = self._books(self.das)['ID19']
        rows = {sheet[f'T{row}'].value: row for row in range(18, 40) if sheet[f'C{row}'].value == 'Titulaire'}
        presence = rows[ID19_NATURES['presence']]
        self.assertEqual(sheet[f'V{presence}'].value, self.row['C1'] + self.row['C4'])
        self.assertEqual((sheet[f'F{presence}'].value, sheet[f'G{presence}'].value), ('BP 100', 'Libreville'))
        self.assertEqual(sheet[f'V{rows[ID19_NATURES["irpp"]]}'].value, self.row['C8'])
        self.assertEqual(sheet[f'V{rows[ID19_NATURES["tcs"]]}'].value, self.row['C7'])
        # avantage logement : V = base, le classeur calcule AG = V × 6 % = avantage déclaré
        housing = rows[ID19_NATURES['housing']]
        rate = self.das._l10n_ga_parameter('l10n_ga_aik_housing_rate')
        self.assertTrue(self.row['aik_housing'])
        self.assertAlmostEqual(sheet[f'V{housing}'].value * rate, self.row['aik_housing'], delta=1)
        self.assertEqual(self.row['aik_housing'] + self.row['aik_utilities'] + self.row['aik_domestic'], self.row['C2'])

    def test_snapshot_and_download(self):
        self.das.with_user(self.declarant).action_validate()
        frozen = self.das.snapshot_attachment_ids.filtered(lambda a: a.name.endswith('.xlsm'))
        self.assertEqual(sorted(frozen.mapped('name')), [f'NIF-TEST-ID19-{YEAR}.xlsm', f'NIF-TEST-ID21-{YEAR}.xlsm'])
        action = self.das.action_download_official_workbooks()
        attachment = self.env['ir.attachment'].browse(int(action['url'].split('/')[3].split('?')[0]))
        self.assertEqual(attachment.mimetype, 'application/zip')
        with zipfile.ZipFile(io.BytesIO(base64.b64decode(attachment.datas))) as archive:
            self.assertEqual(sorted(archive.namelist()), sorted(frozen.mapped('name')))

    def test_address_and_capacity_warnings(self):
        codes = self.das.issue_ids.mapped('code')
        self.assertNotIn('GA_DAS_ID19_ADDRESS', codes)
        self.employee.version_id.private_city = False
        self.das.action_compute()
        self.assertIn('GA_DAS_ID19_ADDRESS', self.das.issue_ids.mapped('code'))
        generator = self.das._generator()
        self.assertFalse(generator._capacity_issues(self.das))

    def test_no_official_workbook_for_other_types(self):
        id10 = self._find(self._type('ID10'), MONTHS[0])
        self.assertFalse(id10.has_official_workbooks)
        self.assertEqual(id10._l10n_ga_official_workbooks(), [])
        with self.assertRaisesRegex(Exception, 'Aucun classeur officiel'):
            id10.action_download_official_workbooks()
