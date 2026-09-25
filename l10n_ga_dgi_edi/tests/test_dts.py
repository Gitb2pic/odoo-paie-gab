import base64
import io
from datetime import date

from openpyxl import load_workbook

from odoo.tests import tagged

from .common import F16_CNAMGS, F16_CNSS, F16_SOCIAL_BASE, GaDeclarationCase

Q3 = (date(2026, 7, 1), date(2026, 9, 30))
MONTHS = [(date(2026, month, 1), date(2026, month, day)) for month, day in ((7, 31), (8, 31), (9, 30))]
CNSS_EMPLOYER = ('GA_CNSS_PF', 'GA_CNSS_AT', 'GA_CNSS_AVID')


@tagged('post_install', '-at_install')
class TestDts(GaDeclarationCase):
    """DTS CNSS et CNAMGS trimestrielles (base 03 §3, §5 ; plan 4.3)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.cnss, cls.cnamgs = cls._type('DTS_CNSS'), cls._type('DTS_CNAMGS')

    def _dts(self, decl_type):
        declaration = self._find(decl_type, Q3)
        self.assertEqual(len(declaration), 1, 'DTS préparée à la validation des bulletins (Observer)')
        return declaration

    @staticmethod
    def _sum(slips, *codes):
        return sum(line.total for line in slips.line_ids if line.code in codes)

    def _quarter(self, employee):
        return self.env['hr.payslip'].concat(*(self._f16_slip(employee, period) for period in MONTHS))

    def test_types(self):
        self.assertEqual(
            (self.cnss.periodicity, self.cnss.period_basis, self.cnss.authority), ('quarterly', 'period', 'cnss')
        )
        self.assertEqual(self.cnss._due_date(Q3[1]), date(2026, 10, 30))
        self.assertEqual(self.cnamgs._due_date(date(2026, 12, 31)), date(2027, 1, 30))
        self.assertEqual(self.cnamgs.authority, 'cnamgs')

    def test_full_quarter_cnss(self):
        employee = self._f16_employee()
        slips = self._quarter(employee)
        declaration = self._dts(self.cnss)
        self.assertEqual(declaration.name, 'DTS_CNSS T3 2026')
        self.assertEqual(declaration.due_date, date(2026, 10, 30))
        values = self._values(declaration)
        self.assertEqual((values['HDR_NUMBER'], values['HDR_YEAR'], values['HDR_QUARTER']), ('CNSS-EMP', 2026, 3))
        self.assertEqual(values['EMPLOYEES'], 1)
        self.assertEqual((values['GROSS'], values['BASE']), (3 * F16_SOCIAL_BASE, 3 * F16_SOCIAL_BASE))
        # Totaux DTS = Σ des lignes CNSS des bulletins du trimestre.
        self.assertEqual(values['SAL'], 3 * F16_CNSS)
        self.assertEqual(values['SAL'], -self._sum(slips, 'GA_CNSS_SAL'))
        self.assertEqual([values[code] for code in ('PF', 'AT', 'AVID')], [self._sum(slips, c) for c in CNSS_EMPLOYER])
        self.assertEqual(values['TOTAL'], -self._sum(slips, 'GA_CNSS_SAL') + self._sum(slips, *CNSS_EMPLOYER))
        self.assertEqual(declaration.amount_total, values['TOTAL'])
        self.assertFalse(declaration.issue_ids)
        detail = declaration.detail_ids
        self.assertEqual(len(detail), 1)
        payload = detail.payload
        self.assertEqual((payload['number'], payload['name'], payload['hire_date']), ('CNSS-F16', 'F16', '2025-01-01'))
        self.assertEqual([payload[f'GROSS_m{i}'] for i in (1, 2, 3)], [F16_SOCIAL_BASE] * 3)
        self.assertEqual(detail.amount, values['TOTAL'])
        self.assertEqual(detail.payslip_line_ids.slip_id, slips)

    def test_full_quarter_cnamgs(self):
        slips = self._quarter(self._f16_employee())
        declaration = self._dts(self.cnamgs)
        values = self._values(declaration)
        self.assertEqual(values['SAL'], 3 * F16_CNAMGS)
        self.assertEqual(values['PAT'], self._sum(slips, 'GA_CNAMGS_PAT'))
        self.assertEqual(values['TOTAL'], -self._sum(slips, 'GA_CNAMGS_SAL') + self._sum(slips, 'GA_CNAMGS_PAT'))
        self.assertEqual(declaration.detail_ids.payload['number'], 'CNAMGS-F16')
        self.assertFalse(declaration.issue_ids)

    def test_entry_and_exit_during_quarter(self):
        arrival = self._f16_employee('Arrivée', start=date(2026, 8, 15))
        for period in MONTHS[1:]:
            self._f16_slip(arrival, period)
        leaver = self._f16_employee('Départ', contract_date_end=date(2026, 8, 31))
        for period in MONTHS[:2]:
            self._f16_slip(leaver, period)
        declaration = self._dts(self.cnss)
        rows = {detail.label: detail.payload for detail in declaration.detail_ids}
        self.assertEqual(rows['Arrivée']['hire_date'], '2026-08-15')
        self.assertEqual(rows['Arrivée']['GROSS_m1'], 0)
        self.assertTrue(0 < rows['Arrivée']['GROSS_m2'] < rows['Arrivée']['GROSS_m3'], 'août proratisé')
        self.assertEqual(rows['Départ']['departure_date'], '2026-08-31')
        self.assertEqual(rows['Départ']['GROSS_m3'], 0)
        self.assertEqual(self._values(declaration)['EMPLOYEES'], 2)
        self.assertEqual(sum(declaration.detail_ids.mapped('amount')), declaration.amount_total)

    def test_ceilings(self):
        cadre = self._employee('Cadre', 2_000_000, ssnid='C1', l10n_ga_cnamgs_number='M1', l10n_ga_nif='N1')
        dirigeant = self._employee('Dirigeant', 3_000_000, ssnid='C2', l10n_ga_cnamgs_number='M2', l10n_ga_nif='N2')
        for employee in (cadre, dirigeant):
            self._payslip(employee, *MONTHS[2]).action_payslip_done()
        cnss = {d.label: d.payload for d in self._dts(self.cnss).detail_ids}
        cnamgs = {d.label: d.payload for d in self._dts(self.cnamgs).detail_ids}
        self.assertEqual((cnss['Cadre']['BASE_m3'], cnss['Dirigeant']['BASE_m3']), (1_500_000, 1_500_000))
        self.assertEqual((cnamgs['Cadre']['BASE_m3'], cnamgs['Dirigeant']['BASE_m3']), (2_000_000, 2_500_000))
        self.assertEqual(cnss['Dirigeant']['GROSS_m3'], 3_000_000)
        self.assertFalse(self._dts(self.cnss).issue_ids)

    def test_two_payslips_same_month_over_ceiling(self):
        cadre = self._employee('Cadre', 1_400_000, ssnid='C1', l10n_ga_cnamgs_number='M1', l10n_ga_nif='N1')
        self._payslip(cadre, *MONTHS[2]).action_payslip_done()
        self._payslip(cadre, *MONTHS[2], inputs={'GA_13M': 1_000_000}).action_payslip_done()
        declaration = self._dts(self.cnss)
        issue = declaration.issue_ids.filtered(lambda i: i.code == 'GA_DTS_CEILING')
        self.assertEqual((issue.severity, issue.employee_id), ('warning', cadre))
        self.assertIn('09/2026', issue.message)

    def test_missing_numbers(self):
        employee = self._f16_employee('Sans numéros', ssnid=False, l10n_ga_cnamgs_number=False)
        self._f16_slip(employee, MONTHS[0])
        self.assertEqual(self._dts(self.cnamgs).issue_ids.mapped('code'), ['GA_DECL_NO_CNAMGS'])
        self.assertEqual(self._dts(self.cnss).issue_ids.mapped('code'), ['GA_DECL_NO_CNSS'])
        employee.l10n_ga_cnamgs_number = 'CNAMGS-OK'  # correction après validation du bulletin
        cnamgs = self._dts(self.cnamgs)
        cnamgs.action_compute()
        self.assertFalse(cnamgs.issue_ids)
        self.assertEqual(cnamgs.detail_ids.payload['number'], 'CNAMGS-OK')

    def test_payment_date_does_not_move_quarter(self):
        self._f16_slip(self._f16_employee(), MONTHS[2], payment_date=date(2026, 10, 5))
        self.assertEqual(self._values(self._dts(self.cnss))['SAL'], F16_CNSS)

    def test_excel_and_pdf(self):
        self._quarter(self._f16_employee())
        declaration = self._dts(self.cnss)
        declaration.with_user(self.declarant).action_validate()
        xlsx = declaration.snapshot_attachment_ids.filtered(lambda a: a.name.endswith('.xlsx'))
        book = load_workbook(io.BytesIO(base64.b64decode(xlsx.datas)))
        rows = [[cell.value for cell in row] for row in book['Détails'].iter_rows()]
        headers = rows[0]
        self.assertEqual(headers[:4], ['N° CNSS', 'Nom et prénoms', 'Date d’entrée', 'Date de sortie'])
        self.assertIn('Salaires soumis 07/2026', headers)
        record = dict(zip(headers, rows[1], strict=True))
        self.assertEqual((record['N° CNSS'], record['Salaires soumis 09/2026']), ('CNSS-F16', F16_SOCIAL_BASE))
        self.assertEqual(record['Cotisation salariale (pensions)'], 3 * F16_CNSS)
        html = base64.b64decode((declaration.snapshot_attachment_ids - xlsx).datas).decode()
        self.assertIn('Détail nominatif', html)
        self.assertIn('CNSS-F16', html)
