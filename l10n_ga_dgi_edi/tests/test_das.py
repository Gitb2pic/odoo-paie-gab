import base64
import io
from ast import literal_eval
from datetime import date

from openpyxl import load_workbook

from odoo.tests import tagged

from .common import GaDeclarationCase

YEAR = 2026
MONTHS = [(date(YEAR, m, 1), date(YEAR, m, d)) for m, d in ((1, 31), (2, 28), (3, 31))]
DAS_PERIOD = (date(YEAR, 1, 1), date(YEAR, 12, 31))
TAX_CODES = {'C7': ('GA_TCS',), 'C8': ('GA_IRPP', 'GA_IRPP_REGUL'), 'C9': ('GA_CFP',)}
ID10_BOXES = {'C7': 'L41', 'C8': 'L40', 'C9': 'R56', 'C10': 'L42'}


@tagged('post_install', '-at_install')
class TestDas(GaDeclarationCase):
    """DAS ID19 à ID22 (base 06 §3 et §4, plan 4.4)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.das_type, cls.id10 = cls._type('DAS'), cls._type('ID10')
        cls.employee = cls._f16_employee('Titulaire', l10n_ga_job_code='E01', l10n_ga_level_code='N1')
        cls.slips = cls.env['hr.payslip'].concat(*(cls._f16_slip(cls.employee, period) for period in MONTHS))

    def _das(self):
        wizard = self.env['l10n_ga.das.wizard'].create({'company_id': self.company.id, 'year': YEAR})
        action = wizard.action_prepare()
        declaration = self.env['l10n_ga.declaration'].browse(action['res_id'])
        self.assertEqual(
            (declaration.type_id, declaration.date_from, declaration.date_to), (self.das_type, *DAS_PERIOD)
        )
        return declaration

    @staticmethod
    def _row(declaration, employee):
        return declaration.detail_ids.filtered(lambda d: d.employee_id == employee).payload

    @staticmethod
    def _sum(slips, *codes):
        return sum(line.total for line in slips.line_ids if line.code in codes)

    @staticmethod
    def _presence(slips):
        """Part imposable des rubriques classées en colonne (1) (catalogue, F16)."""
        return sum(
            line.total - line.l10n_ga_tax_exempt
            for line in slips.line_ids
            if line.salary_rule_id.l10n_ga_das_column == 'presence'
        )

    def _codes(self, declaration):
        return set(declaration.issue_ids.mapped('code'))

    def test_type_and_preparation(self):
        self.assertEqual((self.das_type.periodicity, self.das_type.period_basis), ('yearly', 'payment_date'))
        self.assertEqual(self.das_type._due_date(DAS_PERIOD[1]), date(YEAR + 1, 4, 30))
        self.assertFalse(self._find(self.das_type, DAS_PERIOD), 'jamais préparée à chaque bulletin (D-93)')
        self.env['l10n_ga.declaration']._cron_prepare_due_declarations(today=date(YEAR + 1, 4, 22))
        declaration = self._find(self.das_type, DAS_PERIOD)
        self.assertEqual((declaration.state, declaration.name), ('computed', f'DAS {YEAR}'))
        self.assertEqual(self._das(), declaration, 'l’assistant reprend la DAS existante')

    def test_das_equals_sum_of_id10(self):
        declaration = self._das()
        values = self._values(declaration)
        id10s = self.env['l10n_ga.declaration'].search(
            [('type_id', '=', self.id10.id), ('company_id', '=', self.company.id), ('date_from', '>=', DAS_PERIOD[0])]
        )
        self.assertEqual(len(id10s), 3)
        for das_box, id10_box in ID10_BOXES.items():
            with self.subTest(box=das_box):
                self.assertEqual(values[das_box], sum(self._values(d)[id10_box] for d in id10s))
        self.assertEqual(values['C7'], -self._sum(self.slips, 'GA_TCS'))
        self.assertEqual(values['C8'], -self._sum(self.slips, *TAX_CODES['C8']))
        self.assertEqual(values['C9'], self._sum(self.slips, 'GA_CFP'))
        self.assertEqual(values['C11'], values['C7'] + values['C8'] + values['C9'] + values['C10'])
        self.assertNotIn('GA_DAS_ID10', self._codes(declaration))
        self.assertIn('GA_DAS_PAYMENTS', self._codes(declaration), 'aucune quittance : note explicative ID22')
        self.assertEqual(values['EMPLOYEES'], 1)

    def test_columns_from_das_classification(self):
        declaration = self._das()
        row = self._row(declaration, self.employee)
        presence = self._presence(self.slips)
        contributions = -self._sum(self.slips, 'GA_CNSS_SAL', 'GA_CNAMGS_SAL')
        self.assertEqual(row['C1'], presence - contributions, '(1) nette des cotisations (point 09-6)')
        self.assertEqual(row['C6'], sum(row[c] for c in ('C1', 'C2', 'C3', 'C4', 'C5')))
        transport = sum(
            line.l10n_ga_tax_exempt
            for line in self.slips.line_ids
            if line.salary_rule_id.l10n_ga_das_exempt_column == 'nt_transport'
        )
        self.assertEqual(row['NT_TRANSPORT'], transport)
        self.assertEqual(
            row['NT_TOTAL'], sum(row[c] for c in ('NT_HOUSING', 'NT_TRANSPORT', 'NT_DOMESTIC', 'NT_OTHER'))
        )
        self.assertEqual(
            (row['number'], row['nif'], row['job_code'], row['months']), ('CNSS-Titulaire', 'NIF-Titulaire', 'E01', 3)
        )
        self.assertEqual(self._values(declaration)['C1'], row['C1'], 'ID21 = Σ des lignes arrondies')
        self.company.l10n_ga_das_presence_net = False
        declaration.action_compute()
        self.assertEqual(self._row(declaration, self.employee)['C1'], presence, '(1) brute sur option')

    def test_employee_leaving_during_year(self):
        leaver = self._f16_employee(
            'Sortant', contract_date_end=date(YEAR, 2, 28), l10n_ga_job_code='E', l10n_ga_level_code='N'
        )
        for period in MONTHS[:2]:
            self._f16_slip(leaver, period)
        row = self._row(self._das(), leaver)
        self.assertEqual((row['period'], row['months']), ('01/01 – 28/02', 2))

    def test_switch_year_with_opening(self):
        joiner = self._f16_employee('Repris', l10n_ga_job_code='E', l10n_ga_level_code='N')
        opening = self.env['l10n_ga.ytd.opening'].create(
            {
                'employee_id': joiner.id,
                'company_id': self.company.id,
                'year': YEAR,
                'date_from': date(YEAR, 1, 1),
                'date_to': date(YEAR, 2, 28),
                'taxable': 1_000_000,
                'benefits': 100_000,
                'contributions': 70_000,
                'irpp': 40_000,
                'tcs': 25_000,
                'fnh': 20_000,
                'tax_exempt': 30_000,
            }
        )
        march = self._f16_slip(joiner, MONTHS[2])
        declaration = self._das()
        row = self._row(declaration, joiner)
        self.assertEqual(row['months'], 3, 'janvier-février (cumuls) + mars')
        self.assertEqual(row['C8'], opening.irpp - self._sum(march, *TAX_CODES['C8']))
        march_c1 = self._presence(march) + self._sum(march, 'GA_CNSS_SAL', 'GA_CNAMGS_SAL')
        self.assertEqual(row['C1'], opening.taxable - opening.benefits - opening.contributions + march_c1)
        self.assertGreaterEqual(row['C2'], opening.benefits)
        self.assertGreaterEqual(row['NT_OTHER'], opening.tax_exempt)
        codes = self._codes(declaration)
        self.assertNotIn('GA_DAS_ID10', codes, 'mois repris des cumuls comptés dans le rapprochement')
        self.assertIn('GA_DAS_OPENING_CFP', codes)

    def test_id22_from_receipts(self):
        january = self._find(self.id10, MONTHS[0])
        january.with_user(self.declarant).action_validate()
        january.with_user(self.declarant).action_mark_filed()
        values = self._values(january)
        Payment = self.env['l10n_ga.declaration.payment'].with_user(self.declarant)
        Payment.create(
            {
                'declaration_id': january.id,
                'kind': 'rs',
                'amount': values['L40'] + values['L41'],
                'receipt_number': 'Q-RS-01',
            }
        )
        Payment.create(
            {'declaration_id': january.id, 'kind': 'fnh', 'amount': values['L42'], 'receipt_number': 'Q-FNH-01'}
        )
        Payment.create(
            {'declaration_id': january.id, 'kind': 'cfp', 'amount': values['R56'], 'receipt_number': 'Q-CFP-01'}
        )
        self.assertEqual(january.state, 'paid')
        declaration = self._das()
        das = self._values(declaration)
        self.assertEqual(
            (das['PAID_RS'], das['PAID_FNH'], das['PAID_CFP']),
            (values['L40'] + values['L41'], values['L42'], values['R56']),
        )
        self.assertEqual(das['PAID_TOTAL'], january.amount_paid)
        receipts = declaration.detail_ids.filtered(lambda d: not d.employee_id)
        self.assertEqual(sorted(receipts.mapped('label')), ['Q-CFP-01', 'Q-FNH-01', 'Q-RS-01'])
        self.assertEqual(receipts[0].payload['month'], '01/2026')
        self.assertFalse(declaration.issue_ids.filtered(lambda i: i.code == 'GA_DECL_TOTALS'))

    def test_id20_tranches_and_id19(self):
        small = self._employee(
            'Petit', 60_000, ssnid='C-P', l10n_ga_nif='N-P', l10n_ga_job_code='E', l10n_ga_level_code='N'
        )
        big = self._employee(
            'Cadre', 1_800_000, ssnid='C-C', l10n_ga_nif='N-C', l10n_ga_job_code='E', l10n_ga_level_code='N'
        )
        for employee in (small, big):
            self._payslip(employee, *MONTHS[0]).action_payslip_done()
        declaration = self._das()
        rows = {d.employee_id: d.payload for d in declaration.detail_ids if d.employee_id}
        self.assertEqual((rows[small]['id19'], rows[small]['tranche']), (False, 'B'))
        self.assertEqual((rows[big]['id19'], rows[big]['tranche']), (True, 'A'))
        self.assertEqual((rows[self.employee]['id19'], rows[self.employee]['tranche']), (True, 'B'))
        values = self._values(declaration)
        self.assertEqual((values['ID20_A_COUNT'], values['ID20_B_COUNT'], values['ID19_COUNT']), (1, 2, 2))
        self.assertEqual(values['ID20_A_TOTAL'] + values['ID20_B_TOTAL'], values['C6'])

    def test_employee_checks(self):
        self._f16_slip(self._f16_employee('Sans NIF', l10n_ga_nif=False), MONTHS[0])
        codes = self._codes(self._das())
        self.assertIn('GA_DAS_NO_NIF', codes)
        self.assertIn('GA_DAS_NO_JOB_CODE', codes)

    def test_missing_id10_and_mismatch(self):
        self._find(self.id10, MONTHS[1]).action_cancel()
        codes = self._codes(self._das())
        self.assertIn('GA_DAS_ID10_MISSING', codes)
        self.assertIn('GA_DAS_ID10', codes, 'février absent des ID10 : totaux différents')

    def test_workbook_and_pdf(self):
        second = self._f16_employee('Second', l10n_ga_job_code='E', l10n_ga_level_code='N')
        self._f16_slip(second, MONTHS[0])
        self.env.ref('l10n_ga_hr_payroll.rule_parameter_das_id21_lines_20000101').parameter_value = '1'
        declaration = self._das()
        declaration.with_user(self.declarant).action_validate()
        xlsx = declaration.snapshot_attachment_ids.filtered(lambda a: a.name.endswith('.xlsx'))
        book = load_workbook(io.BytesIO(base64.b64decode(xlsx.datas)))
        self.assertEqual(book.sheetnames[:4], ['ID20', 'ID21 (1)', 'ID21 (2)', 'ID22'])
        self.assertEqual(len([n for n in book.sheetnames if n.startswith('ID19')]), 2)
        texts = [str(c.value) for row in book['ID22'].iter_rows() for c in row if c.value is not None]
        self.assertIn('Récapitulatif des feuilles de l’ID21', texts)
        html = base64.b64decode((declaration.snapshot_attachment_ids - xlsx).datas).decode()
        for title in (
            'ID20 — État de la masse salariale',
            'ID21 — Bordereau détaillé',
            'ID22 — Bordereau récapitulatif',
            'ID19 — Bulletin individuel',
        ):
            self.assertIn(title, html)

    def test_check_screen(self):
        wizard = self.env['l10n_ga.das.wizard'].create({'company_id': self.company.id, 'year': YEAR})
        action = wizard.action_open_checks()
        declaration = self._find(self.das_type, DAS_PERIOD)
        self.assertEqual(action['domain'], [('declaration_id', '=', declaration.id)])
        screen = self.env.ref('l10n_ga_dgi_edi.l10n_ga_check_issue_action_das')
        issues = self.env['l10n_ga.check.issue'].search(literal_eval(screen.domain))
        self.assertEqual(issues, declaration.issue_ids)
