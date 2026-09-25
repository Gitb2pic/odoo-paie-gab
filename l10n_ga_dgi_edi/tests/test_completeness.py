"""Complétude C-3 : cas limites du moteur non couverts par les tests des étapes 4.1 à 4.4."""

import io
from datetime import date, datetime
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from odoo.fields import Command
from odoo.tests import tagged

from ..models.declaration_generator import L10nGaDeclarationGenerator
from ..renderers.xlsm_template import XlsmTemplateRenderer
from ..renderers.xlsx_builder import XlsxDeclarationBuilder
from ..renderers.xlsx_html import workbook_to_html
from .common import F16_IRPP, SEPT, FakeGenerator, GaDeclarationCase
from .test_xlsm_macros import VBA_PROJECT, make_xlsm, sheet_xml


def _book_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@tagged('post_install', '-at_install')
class TestEngineCompleteness(GaDeclarationCase):
    # --- rendu PDF = classeur (xlsx_html) -------------------------------------------------------

    def test_excel_to_html_formats(self):
        book = Workbook()
        sheet = book.active
        sheet.merge_cells('A1:C1')
        sheet['A1'] = 'Titre fusionné'
        sheet['A1'].font = Font(bold=True, color='FFFF0000')
        sheet['A2'], sheet['B2'], sheet['C2'] = 'Taux', 'masquée', 0.005
        sheet['C2'].number_format = '0.0%'
        sheet.column_dimensions['B'].hidden = True
        sheet['A3'] = 'ligne masquée'
        sheet.row_dimensions[3].hidden = True
        sheet['A4'] = date(2026, 9, 30)
        sheet['C4'] = '=1+1'
        sheet['A5'] = 1234567
        sheet['A5'].number_format = '#,##0'
        sheet['A5'].fill = PatternFill('solid', fgColor='FFFFFF00')
        sheet['A6'] = 'Un texte long qui déborde jusqu’à la cellule remplie'
        sheet['C6'] = 'fin'
        sheet.freeze_panes = 'A2'
        html = workbook_to_html(_book_bytes(book), 700)
        self.assertIn('colspan="2"', html)  # A1:C1 sans la colonne B masquée
        self.assertIn('font-weight:bold', html)
        self.assertIn('color:#FF0000', html)
        self.assertIn('0,5 %', html)
        self.assertNotIn('masquée', html)
        self.assertIn('30/09/2026', html)
        self.assertNotIn('1+1', html)  # jamais de formule affichée
        self.assertIn('1 234 567', html)
        self.assertIn('background-color:#FFFF00', html)
        self.assertIn('<thead', html)  # ligne figée répétée sur chaque page
        self.assertIn('overflow:hidden;white-space:nowrap;', html)  # débordement arrêté avant « fin »

    def test_sheets_on_separate_pages(self):
        book = Workbook()
        book.active['A1'] = 'page 1'
        book.create_sheet('Deux')['A1'] = 'page 2'
        hidden = book.create_sheet('Cachée')
        hidden['A1'] = 'jamais imprimée'
        hidden.sheet_state = 'hidden'
        html = workbook_to_html(_book_bytes(book), 700)
        self.assertEqual(html.count('page-break-before:always'), 1)
        self.assertNotIn('jamais imprimée', html)

    # --- gabarits : listes, feuillets, formules -------------------------------------------------

    def test_template_rows_sheets_and_formulas(self):
        renderer = XlsmTemplateRenderer(make_xlsm())
        self.assertEqual(renderer.sheet_names(), ['ID', 'Annexe'])
        copy = renderer.copy_sheet('Annexe', 'Annexe (2)')
        renderer.rows('ID', 20, [('A', 1), ('B', 2)], ('A', 'B'), capacity=4)
        renderer.remove_sheet(copy)
        content = renderer.build()
        files = sheet_xml(content)
        self.assertEqual(files['xl/vbaProject.bin'], VBA_PROJECT)
        for name, data in files.items():
            if name.startswith('xl/worksheets/'):
                self.assertNotIn(b'<f>', data)  # la formule =1+1 du gabarit (C11) est retirée
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book.sheetnames, ['ID', 'Annexe'])
        sheet = book['ID']
        self.assertEqual([sheet['A20'].value, sheet['B21'].value], ['A', 2])
        self.assertIsNone(sheet['A23'].value)
        self.assertIsNone(sheet['C11'].value)

    # --- moteur ---------------------------------------------------------------------------------

    def test_date_box_and_builder(self):
        decl_type = self._declaration_type('T_DATE')
        decl_type.box_ids = [Command.create({'code': 'D_DEPOT', 'name': 'Date de dépôt', 'value_kind': 'date'})]
        fake = FakeGenerator(values={'D_DEPOT': date(2026, 10, 15)})
        with patch.object(L10nGaDeclarationGenerator, '_get', lambda registry, key: fake):
            declaration = self._declaration(decl_type=decl_type)
            declaration.action_compute()
            content, _ext = declaration._render_xlsx()
        line = self._box(declaration, 'D_DEPOT')
        self.assertEqual((line.value_date, line._value()), (date(2026, 10, 15), date(2026, 10, 15)))
        rows = [[c.value for c in row] for row in load_workbook(io.BytesIO(content))['Déclaration'].iter_rows()]
        self.assertIn(['D_DEPOT', 'Date de dépôt', datetime(2026, 10, 15)], rows)
        builder = XlsxDeclarationBuilder(title='T')
        builder.boxes(['Case'], [('X', True)])  # booléen écrit en texte, jamais en nombre
        self.assertTrue(load_workbook(io.BytesIO(builder.build()))['Déclaration'])

    def test_observer_on_payslip_cancel(self):
        auto = self._declaration_type('T_AUTO_CANCEL', auto_create=True)
        slip = self._f16_slip()
        declaration = self._find(auto)
        self.assertEqual(self._box(declaration, 'B_IRPP').value_amount, F16_IRPP)
        slip.action_payslip_cancel()
        self.assertEqual(self._box(declaration, 'B_IRPP').value_amount, 0, 'bulletin annulé : sorti de la déclaration')
        self.assertEqual(SEPT[1].month, 9)

    def test_issue_opens_record_to_fix(self):
        employee = self._f16_employee('Sans numéro', ssnid=False)
        self._f16_slip(employee)
        declaration = self._declaration()
        declaration.action_compute()
        issue = declaration.issue_ids.filtered(lambda i: i.code == 'GA_DECL_NO_CNSS')
        action = issue.action_open_record()
        self.assertEqual((action['res_model'], action['res_id']), ('hr.employee', employee.id))
