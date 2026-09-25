import base64
import io

from openpyxl import load_workbook

from odoo.tests import tagged

from ..renderers.xlsm_template import XlsmTemplateRenderer
from ..renderers.xlsx_html import FONT_STACK, workbook_to_html
from .common import F16_IRPP, GaDeclarationCase


@tagged('post_install', '-at_install')
class TestReportLayout(GaDeclarationCase):
    """PDF = rendu fidèle du classeur Excel de la déclaration (mise en page de l'imprimé officiel)."""

    def test_official_form_rendered_like_excel(self):
        self._f16_slip()
        declaration = self._find(self._type('ID10'))
        html = str(declaration._l10n_ga_excel_html(740))
        self.assertIn(FONT_STACK, html)
        self.assertIn('REPUBLIQUE GABONAISE', html)
        self.assertIn('Montant global dû :', html)
        self.assertIn('23 195', html)  # montant au format du gabarit (# ##0)
        self.assertIn('0,5 %', html)  # taux au format pourcentage du gabarit
        self.assertIn('colspan="6"', html)  # cellules fusionnées P40:U40
        self.assertIn('border-', html)
        self.assertNotIn('=SUM', html)
        declaration.with_user(self.declarant).action_validate()
        report = declaration.snapshot_attachment_ids.filtered(lambda a: not a.name.endswith('.xlsx'))
        self.assertIn('23 195', base64.b64decode(report.datas).decode())

    def test_values_take_label_font(self):
        template = self._type('ID28')._template_bytes()
        renderer = XlsmTemplateRenderer(template)
        renderer.boxes([('D22', 'Société'), ('Q43', 2_775)])
        sheet = load_workbook(io.BytesIO(renderer.build()))['CFP']
        self.assertEqual(sheet['D22'].font.sz, sheet['A22'].font.sz)  # 14 comme « Raison sociale : »
        self.assertGreaterEqual(sheet['Q43'].font.sz, 11)

    def test_new_workbook_layout(self):
        self._f16_slip()
        declaration = self._find(self._type('ID10'))
        declaration.type_id.template_path = False  # classeur neuf
        content, _extension = declaration._render_xlsx()
        book = load_workbook(io.BytesIO(content))
        sheet = book['Déclaration']
        self.assertEqual(sheet['A1'].font.name, 'Arial')
        self.assertTrue(sheet['A1'].font.b)
        header = next(row for row in sheet.iter_rows() if row[0].value == 'Case')
        self.assertEqual(header[0].fill.fgColor.rgb, 'FFD9D9D9')
        self.assertNotIn('Année', [row[0].value for row in sheet.iter_rows()])  # pas d'en-tête d'ID10 ici
        self.assertEqual(book['Détails'].freeze_panes, 'A4')  # titre + en-tête figés et répétés
        html = workbook_to_html(content, 1068)
        self.assertIn('<thead ', html)  # en-tête de l'état nominatif répété sur chaque page
        self.assertIn(f'{F16_IRPP:,}'.replace(',', ' '), html)
