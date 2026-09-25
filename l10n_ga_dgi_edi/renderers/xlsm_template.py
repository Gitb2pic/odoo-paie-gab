"""Gabarit officiel rempli (``openpyxl``, ``keep_vba=True``) : le projet VBA des ``.xlsm`` est conservé (F10).

Seules des valeurs sont écrites : une chaîne qui commence par « = » est forcée en texte, pour
qu'aucune formule ne soit introduite (règle d'or 10). Les formules déjà présentes dans une cellule
remplie sont remplacées par la valeur.
"""

import io

from openpyxl import load_workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.styles import Font
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.exceptions import CellCoordinatesException

from .base import DeclarationBuilder


class XlsmTemplateRenderer(DeclarationBuilder):
    def __init__(self, template_bytes):
        self._book = load_workbook(io.BytesIO(template_bytes), keep_vba=True)

    def _cell(self, ref):
        """``Feuille!P40`` ou ``P40`` (première feuille)."""
        if '!' in ref:
            sheet_name, coordinate = ref.rsplit('!', 1)
            sheet = self._book[sheet_name.strip("'")]
        else:
            sheet, coordinate = self._book.worksheets[0], ref
        try:
            coordinate_from_string(coordinate)
        except CellCoordinatesException as error:
            raise ValueError(f'Cellule invalide : {ref}') from error
        return sheet[coordinate]

    @staticmethod
    def _harmonize_font(cell):
        """Valeur saisie à la taille du libellé de sa ligne (les zones de saisie des gabarits DGI sont
        souvent en Calibri 11 sous des libellés en 14 ou 16) : lisible à l'écran comme sur le PDF."""
        labels = [
            other
            for other in cell.parent[cell.row]
            if other.column < cell.column and isinstance(other.value, str) and other.value.strip()
        ]
        if not labels:
            return
        label = labels[-1].font
        if (label.sz or 0) > (cell.font.sz or 0):
            font = cell.font
            cell.font = Font(name=label.name, sz=label.sz, b=font.b, i=font.i, u=font.u, color=font.color)

    def _write(self, cell, value):
        if value is None or value is False:
            cell.value = None
            return
        cell.value = value
        self._harmonize_font(cell)
        if isinstance(value, str) and value.startswith('='):
            cell.data_type = TYPE_STRING  # valeur, jamais formule

    def header(self, values):
        for ref, value in values.items():
            self._write(self._cell(ref), value)

    def boxes(self, cells):  # pylint: disable=arguments-differ
        """``[(référence de cellule, valeur), ...]`` ou ``(référence, valeur, format numérique)`` : un
        montant écrit dans une cellule au format « Standard » du gabarit reçoit ce format (# ##0)."""
        for ref, value, *number_format in cells:
            cell = self._cell(ref)
            self._write(cell, value)
            if number_format and number_format[0] and cell.number_format == 'General':
                cell.number_format = number_format[0]

    def table(self, sheet, start_row, rows, headers=()):
        """Lignes écrites à partir de ``start_row`` (0 = ligne 1), colonne A, dans la feuille ``sheet``."""
        worksheet = self._book[sheet]
        for offset, row in enumerate([*([list(headers)] if headers else []), *rows]):
            for column, value in enumerate(row, start=1):
                self._write(worksheet.cell(row=start_row + offset + 1, column=column), value)

    def sheet_names(self):
        return self._book.sheetnames

    def copy_sheet(self, name, title):
        """Feuillet supplémentaire : copie d'une feuille du gabarit (avant remplissage)."""
        copy = self._book.copy_worksheet(self._book[name])
        copy.title = title[:31]
        return copy.title

    def remove_sheet(self, name):
        self._book.remove(self._book[name])

    def rows(self, sheet, start_row, rows, columns, capacity=None):
        """Lignes écrites dans les colonnes ``columns`` (lettres) à partir de ``start_row`` (numéro Excel) ;
        les lignes restantes de la zone (``capacity``) sont vidées : aucune formule du gabarit ne subsiste."""
        worksheet = self._book[sheet]
        rows = list(rows)
        for offset in range(max(capacity or 0, len(rows))):
            values = rows[offset] if offset < len(rows) else [None] * len(columns)
            for column, value in zip(columns, values, strict=True):
                self._write(worksheet[f'{column}{start_row + offset}'], value)

    def _strip_formulas(self):
        """Règle d'or 10 : les formules restantes du gabarit (souvent fausses, base 06 §3.3) sont retirées."""
        for worksheet in self._book.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == 'f':
                        cell.value = None

    def build(self):
        self._strip_formulas()
        stream = io.BytesIO()
        self._book.save(stream)
        self.close()
        return stream.getvalue()

    def close(self):
        """Libère la copie de l'archive VBA gardée par ``openpyxl`` (``keep_vba``)."""
        if self._book.vba_archive:
            self._book.vba_archive.close()
            self._book.vba_archive = None
