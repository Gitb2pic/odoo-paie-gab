"""Complétude C-4 : cas limites non couverts par les tests de l'étape 5."""

import base64
import io
from datetime import date

from openpyxl import load_workbook

from odoo.exceptions import ValidationError
from odoo.fields import Command
from odoo.tests import tagged

from .common import SEPT, YEAR, GaWithholdingCase


@tagged('post_install', '-at_install')
class TestCompleteness(GaWithholdingCase):
    @staticmethod
    def _values(declaration):
        return {line.code: line._value() for line in declaration.line_ids}

    def _annual(self, code):
        return self.env['l10n_ga.declaration']._l10n_ga_prepare(self.company, self._type_account(code), *YEAR)

    def test_payment_withholding_lines_never_cumulate(self):
        bill = self._bill(self.provider, 100_000)
        wizard = (
            self.env['account.payment.register']
            .with_context(active_model='account.move', active_ids=bill.ids)
            .create({'payment_date': date(2026, 9, 20)})
        )
        wizard.withholding_line_ids = [
            Command.create(
                {
                    'tax_id': self.ras_20.id,
                    'base_amount': 100_000,
                    'name': 'RAS20-MANUEL',
                    'account_id': wizard.withholding_line_ids[:1].account_id.id,
                }
            )
        ]
        with self.assertRaises(ValidationError):
            wizard._create_payments()

    def test_double_withholding_detected_after_reclassification(self):
        self._pay(self._bill(self.provider, 100_000, day=date(2026, 9, 2)), day=date(2026, 9, 5))
        self.provider.l10n_ga_is_resident = False  # reclassé non-résident en cours de mois
        self.provider.country_id = self.env.ref('base.fr')
        self._pay(self._bill(self.provider, 50_000, day=date(2026, 9, 10)), day=date(2026, 9, 15))
        id18, id27 = self._monthly('ID18'), self._monthly('ID27')
        (id18 | id27).action_compute()
        for declaration in (id18, id27):
            with self.subTest(declaration=declaration.name):
                self.assertIn('GA_RAS_DOUBLE', declaration.issue_ids.mapped('code'))

    def test_id27_overflow_copies_sheet(self):
        partners = [
            self._partner(f'Étranger {i:02d}', l10n_ga_is_resident=False, country_id=self.env.ref('base.be').id)
            for i in range(11)
        ]
        for partner in partners:
            self._pay(self._bill(partner, 10_000))
        id27 = self._monthly('ID27')
        id27.with_user(self.declarant).action_validate()
        xlsx = id27.snapshot_attachment_ids.filtered(lambda a: a.name.endswith('.xlsx'))
        book = load_workbook(io.BytesIO(base64.b64decode(xlsx.datas)))
        self.assertEqual(book.sheetnames, ['Bordereau 1', 'Bordereau 1 (2)'])
        first, second = book['Bordereau 1'], book['Bordereau 1 (2)']
        self.assertEqual((first['J43'].value, first['O43'].value), (100_000, 20_000))
        self.assertEqual((second['A33'].value, second['J43'].value), ('Étranger 10', 10_000))
        self.assertIsNone(second['A34'].value)

    def test_foreign_currency_payment(self):
        eur = self.env.ref('base.EUR')
        eur.active = True
        self.env['res.currency.rate'].create(
            {'currency_id': eur.id, 'company_id': self.company.id, 'name': date(2026, 1, 1), 'rate': 1 / 655.957}
        )
        bill = self.env['account.move'].create(
            {
                'move_type': 'in_invoice',
                'partner_id': self.foreigner.id,
                'currency_id': eur.id,
                'invoice_date': date(2026, 9, 5),
                'date': date(2026, 9, 5),
                'invoice_line_ids': [Command.create({'product_id': self.product.id, 'price_unit': 1_000})],
            }
        )
        bill.action_post()
        self._pay(bill)
        values = self._values(self._monthly('ID27'))
        self.assertAlmostEqual(values['TOTAL_BASE'], 655_957, delta=1)
        self.assertAlmostEqual(values['TOTAL_WITHHELD'], 131_191, delta=1)

    def test_annual_paid_amounts_excluding_vat_and_refunds(self):
        vat = self.env['account.tax'].create(
            {'name': 'TVA 18 % achats (test)', 'amount': 18, 'type_tax_use': 'purchase', 'company_id': self.company.id}
        )
        product = self.env['product.product'].create(
            {'name': 'Honoraires', 'type': 'service', 'supplier_taxes_id': [Command.set(vat.ids)]}
        )
        lawyer = self._partner('Maître Obame', l10n_ga_fee_category='C', vat='NIF-AV2')
        bill = self.env['account.move'].create(
            {
                'move_type': 'in_invoice',
                'partner_id': lawyer.id,
                'invoice_date': date(2026, 3, 1),
                'date': date(2026, 3, 1),
                'invoice_line_ids': [Command.create({'product_id': product.id, 'price_unit': 300_000})],
            }
        )
        bill.action_post()
        self.assertEqual(bill.amount_total, 354_000)
        self._pay(bill, day=date(2026, 3, 10))
        self.assertEqual(self._values(self._annual('ID23'))['PAID_OTHER'], 300_000, 'sommes versées hors TVA')

        self._pay(self._bill(self.provider, 100_000, day=date(2026, 1, 5)), day=date(2026, 1, 10))
        refund = self._bill(self.provider, 20_000, day=date(2026, 2, 1), move_type='in_refund')
        self._pay(refund, day=date(2026, 2, 10))
        values = self._values(self._annual('ID26'))
        self.assertEqual((values['TOTAL_PAID'], values['TOTAL_WITHHELD']), (80_000, 7_600))

    def test_cron_prepares_monthly_declarations(self):
        self.env['l10n_ga.declaration']._cron_prepare_due_declarations(today=date(2026, 10, 6))
        for code in ('ID18', 'ID27'):
            with self.subTest(code=code):
                declaration = self._monthly(code)
                self.assertEqual((declaration.state, declaration.due_date), ('computed', date(2026, 10, 15)))
        self.assertEqual(SEPT[1], date(2026, 9, 30))
