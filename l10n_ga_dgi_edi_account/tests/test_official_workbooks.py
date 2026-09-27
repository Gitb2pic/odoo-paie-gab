import io
import warnings
from datetime import date

from openpyxl import load_workbook as _load_workbook

from odoo.tests import tagged
from odoo.tools.misc import file_open

from ..models.generators.id23 import L10nGaDeclarationGeneratorId23
from .common import YEAR, GaWithholdingCase


def load_workbook(*args, **kwargs):
    """Lecture de contrôle : openpyxl ignore les listes déroulantes étendues (x14) — sans effet ici."""
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Data Validation extension')
        return _load_workbook(*args, **kwargs)


@tagged('post_install', '-at_install')
class TestOfficialWorkbooksFees(GaWithholdingCase):
    """D-87, FIX 05 : annexes ID23 et ID26 dans les classeurs officiels de la DGI ; ID24 sans classeur officiel."""

    def _annual(self, code):
        return self.env['l10n_ga.declaration']._l10n_ga_prepare(self.company, self._type_account(code), *YEAR)

    @staticmethod
    def _sheet(declaration):
        ((name, content),) = declaration._l10n_ga_official_workbooks()
        return name, load_workbook(io.BytesIO(content))['SAISIE']

    def test_id26_official_workbook(self):
        self._pay(self._bill(self.provider, 100_000, day=date(2026, 1, 10)), day=date(2026, 1, 20))
        id26 = self._annual('ID26')
        self.assertTrue(id26.has_official_workbooks)
        name, sheet = self._sheet(id26)
        self.assertEqual(name, 'NIF-TEST-ID26-2026.xlsm')
        self.assertEqual((sheet['C9'].value, sheet['C11'].value, sheet['C13'].value), ('NIF-TEST', 2026, 'Annuel'))
        self.assertTrue(sheet['B18'].value.startswith('Prestataire local'))
        self.assertEqual((sheet['C18'].value, sheet['D18'].value), ('NIF-PREST', 100_000))
        self.assertEqual(sheet['E18'].value, '=D18*0.095')  # retenue calculée par le classeur, jamais écrasée

    def test_id23_official_workbook(self):
        lawyer = self._partner('Maître Nze', l10n_ga_fee_category='C', vat='NIF-AV')
        self._pay(self._bill(lawyer, 300_000, day=date(2026, 4, 1)), day=date(2026, 4, 15))
        _name, sheet = self._sheet(self._annual('ID23'))
        self.assertEqual((sheet['B17'].value, sheet['C17'].value), ('Maître Nze', 'NIF-AV'))
        self.assertEqual((sheet['E17'].value, sheet['F17'].value), ('Qualité de non salarié', 300_000))

    def test_quality_labels_exist_in_template(self):
        with file_open('l10n_ga_dgi_edi_account/static/templates/edi-annexe-ID23.xlsm', 'rb') as stream:
            book = load_workbook(io.BytesIO(stream.read()), read_only=True)
        values = {cell for row in book['Referentiel'].iter_rows(values_only=True) for cell in row}
        self.assertLessEqual(set(L10nGaDeclarationGeneratorId23.QUALITY.values()), values)

    def test_id24_has_no_official_workbook(self):
        id24 = self._annual('ID24')
        self.assertFalse(id24.has_official_workbooks)
        self.assertEqual(id24._l10n_ga_official_workbooks(), [])

    def test_capacity_warning(self):
        generator = self.env['l10n_ga.declaration.generator']._get('id26')
        id26 = self._annual('ID26')
        self.assertEqual(generator._capacity_issues(id26, [{}] * 10), [])
        issues = generator._capacity_issues(id26, [{}] * 2001)
        self.assertEqual(issues[0][1], 'GA_RAS_XLSM_CAPACITY')
