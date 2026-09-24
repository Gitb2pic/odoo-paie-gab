"""Import Excel des variables du mois d'un lot (F3, patron 13, D-41 à D-43).

Côté Odoo : modèle du lot (``xlsxwriter``), lecture du classeur (``openpyxl``), référentiels,
aperçu, écriture dans ``hr.payslip.input`` / ``hr.work.entry`` et recalcul. Les filtres de
validation sont purs (``lib/ga_fiscal_core/input_import.py``).
"""

import base64
import io
from datetime import timedelta

import openpyxl
import xlsxwriter

from odoo import fields, models
from odoo.exceptions import UserError

from ..lib.ga_fiscal_core import input_import as ii

LOAN_INPUT_XMLID = 'l10n_ga_hr_payroll.input_type_ga_loan'
MAX_DAY_HOURS = 24  # contrainte standard d'une prestation (C/addons/hr_work_entry/models/hr_work_entry.py:55-59)
KEY_HEADER = 'Matricule'
NAME_HEADER = 'Salarié'
HEADER_SEPARATOR = ' — '
# Largeurs de colonnes du modèle (présentation) : matricule, nom, variables.
KEY_WIDTH, NAME_WIDTH, VALUE_WIDTH = 14, 30, 18


class L10nGaPayslipInputImport(models.TransientModel):
    _name = 'l10n_ga.payslip.input.import'
    _description = 'Import des variables du mois (Gabon)'
    _check_company_auto = True

    payslip_run_id = fields.Many2one(
        'hr.payslip.run', string='Lot de paie', required=True, ondelete='cascade', check_company=True
    )
    company_id = fields.Many2one(related='payslip_run_id.company_id')
    state = fields.Selection(
        [('upload', 'Fichier'), ('preview', 'Aperçu'), ('done', 'Terminé')], default='upload', required=True
    )
    template_file = fields.Binary(string='Modèle', readonly=True, attachment=False)
    template_filename = fields.Char(readonly=True)
    import_file = fields.Binary(string='Fichier à importer', attachment=False)
    import_filename = fields.Char()
    valid_count = fields.Integer(string='Lignes acceptées', readonly=True)
    rejected_count = fields.Integer(string='Lignes refusées', readonly=True)
    issues_text = fields.Text(string='Anomalies', readonly=True)

    # --- référentiels du lot -------------------------------------------------------------------

    def _draft_slips(self):
        self.ensure_one()
        return self.payslip_run_id.slip_ids.filtered(lambda s: s.state == 'draft')

    def _input_types(self):
        """Types d'entrée utilisables sur les structures du lot (``GA_LOAN`` exclu : échéances F1)."""
        self.ensure_one()
        structures = self.payslip_run_id.structure_id | self._draft_slips().struct_id
        country = self.company_id.country_id
        types = self.env['hr.payslip.input.type'].search([('country_id', 'in', [False, country.id])], order='code, id')
        loan_type = self.env.ref(LOAN_INPUT_XMLID)
        return types.filtered(lambda t: t != loan_type and (not t.struct_ids or t.struct_ids & structures))

    def _hours_types(self):
        return self.env['hr.work.entry.type'].search([('l10n_ga_overtime_period', '!=', False)], order='sequence, id')

    def _context(self):
        """``ImportContext`` du pipeline pur."""
        self.ensure_one()
        slips = self._draft_slips()
        columns = {t.code: (ii.AMOUNT, t.id) for t in self._input_types()}
        columns.update({t.code: (ii.HOURS, t.id) for t in self._hours_types()})
        employees = self.env['hr.employee'].search(
            [('company_id', '=', self.company_id.id), ('registration_number', '!=', False)]
        )
        by_name = {}
        for employee in slips.employee_id:
            by_name.setdefault(employee.name, []).append(employee.id)
        return ii.ImportContext(
            columns=columns,
            employees_by_ref={employee.registration_number.strip(): employee.id for employee in employees},
            employees_by_name={name: tuple(ids) for name, ids in by_name.items()},
            batch=frozenset(slips.employee_id.ids),
        )

    # --- modèle Excel --------------------------------------------------------------------------

    def _template_headers(self):
        headers = [KEY_HEADER, NAME_HEADER]
        headers += [f'{t.code}{HEADER_SEPARATOR}{t.name}' for t in self._input_types()]
        headers += [f'{t.code}{HEADER_SEPARATOR}{t.name} (heures)' for t in self._hours_types()]
        return headers

    def _template_bytes(self):
        self.ensure_one()
        stream = io.BytesIO()
        workbook = xlsxwriter.Workbook(stream, {'in_memory': True})
        sheet = workbook.add_worksheet(self.env._('Variables'))
        bold = workbook.add_format({'bold': True, 'text_wrap': True})
        headers = self._template_headers()
        sheet.write_row(0, 0, headers, bold)
        sheet.freeze_panes(1, 2)
        sheet.set_column(0, 0, KEY_WIDTH)
        sheet.set_column(1, 1, NAME_WIDTH)
        sheet.set_column(2, len(headers) - 1, VALUE_WIDTH)
        for row, slip in enumerate(self._draft_slips().sorted(lambda s: s.employee_id.name or ''), start=1):
            sheet.write_string(row, 0, slip.employee_id.registration_number or '')
            sheet.write_string(row, 1, slip.employee_id.name)
        workbook.close()
        return stream.getvalue()

    def action_generate_template(self):
        self.ensure_one()
        self.write(
            {
                'template_file': base64.b64encode(self._template_bytes()),
                'template_filename': f'variables_{self.payslip_run_id.name or self.payslip_run_id.id}.xlsx',
            }
        )
        return self._reopen()

    # --- lecture et pipeline -------------------------------------------------------------------

    def _read_workbook(self):
        """(en-tête, [(n° de ligne, cellules)]) de la première feuille."""
        self.ensure_one()
        if not self.import_file:
            raise UserError(self.env._('Choisissez le fichier Excel à importer.'))
        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.import_file)), read_only=True, data_only=True
            )
        except Exception as error:  # classeur illisible : zip, xml ou format invalide
            raise UserError(self.env._('Fichier Excel illisible : %(error)s', error=error)) from error
        rows = list(workbook.active.iter_rows(values_only=True))
        workbook.close()
        if not rows:
            raise UserError(self.env._('Le fichier est vide.'))
        return rows[0], list(enumerate(rows[1:], start=2))

    def _pipeline(self):
        header, raw_rows = self._read_workbook()
        return ii.run(header, raw_rows, self._context())

    def _issue_message(self, issue):
        employees = self.env['hr.employee'].browse
        messages = {
            ii.MISSING_KEY_COLUMN: lambda: self.env._('Colonne « Matricule » absente : fichier refusé.'),
            ii.DUPLICATE_COLUMN: lambda: self.env._('Colonne %(column)s en double : ignorée.', column=issue.column),
            ii.UNKNOWN_COLUMN: lambda: self.env._(
                'Colonne %(column)s inconnue pour ce lot : ignorée.', column=issue.column
            ),
            ii.UNKNOWN_EMPLOYEE: lambda: self.env._(
                'Salarié introuvable (%(value)s) : ligne refusée.', value=issue.value or '—'
            ),
            ii.NOT_IN_BATCH: lambda: self.env._(
                '%(employee)s n’a pas de bulletin brouillon dans ce lot : ligne refusée.',
                employee=employees(issue.value).name,
            ),
            ii.DUPLICATE_EMPLOYEE: lambda: self.env._(
                '%(employee)s apparaît sur plusieurs lignes : ligne refusée.', employee=employees(issue.value).name
            ),
            ii.INVALID_NUMBER: lambda: self.env._(
                'Colonne %(column)s : « %(value)s » n’est pas un nombre : ligne refusée.',
                column=issue.column,
                value=issue.value,
            ),
            ii.NEGATIVE_NUMBER: lambda: self.env._(
                'Colonne %(column)s : montant négatif %(value)s : ligne refusée.',
                column=issue.column,
                value=issue.value,
            ),
        }
        text = messages[issue.code]()
        return self.env._('Ligne %(row)s : %(text)s', row=issue.row, text=text) if issue.row else text

    def _issues_summary(self, issues):
        return '\n'.join(self._issue_message(issue) for issue in issues)

    def _rejected_rows(self, issues):
        return len({issue.row for issue in issues if issue.row})

    def action_preview(self):
        self.ensure_one()
        rows, issues = self._pipeline()
        self.write(
            {
                'state': 'preview',
                'valid_count': len(rows),
                'rejected_count': self._rejected_rows(issues),
                'issues_text': self._issues_summary(issues) or False,
            }
        )
        return self._reopen()

    # --- écriture ------------------------------------------------------------------------------

    def _write_amount(self, slip, input_type, amount):
        """La valeur importée remplace les saisies non automatiques du type (D-42) ; 0 supprime."""
        manual = slip.input_line_ids.filtered(
            lambda line: (
                line.input_type_id == input_type and not line.l10n_ga_allowance_id and not line.l10n_ga_loan_line_id
            )
        )
        manual.unlink()
        if amount:
            self.env['hr.payslip.input'].create(
                {
                    'payslip_id': slip.id,
                    'input_type_id': input_type.id,
                    'amount': amount,
                    'name': input_type.name,
                    'l10n_ga_imported': True,
                }
            )

    def _write_hours(self, slip, entry_type, hours):
        """Heures décimales réparties sur la période, du dernier jour au premier, ≤ 24 h/jour (D-43)."""
        WorkEntry = self.env['hr.work.entry']
        period = [
            ('employee_id', '=', slip.employee_id.id),
            ('date', '>=', slip.date_from),
            ('date', '<=', slip.date_to),
        ]
        WorkEntry.search(
            [*period, ('work_entry_type_id', '=', entry_type.id), ('l10n_ga_imported', '=', True)]
        ).unlink()
        if not hours:
            return
        used = {}
        blocked = set()
        for entry in WorkEntry.search(period):
            used[entry.date] = used.get(entry.date, 0) + entry.duration
            if entry.state == 'validated':
                blocked.add(entry.date)
        days = []
        day = slip.date_to
        while day >= slip.date_from:
            if day not in blocked:
                days.append((day, MAX_DAY_HOURS - used.get(day, 0)))
            day -= timedelta(days=1)
        WorkEntry.create(
            [
                {
                    'name': f'{entry_type.name} {slip.employee_id.name}',
                    'employee_id': slip.employee_id.id,
                    'version_id': slip.version_id.id,
                    'company_id': slip.company_id.id,
                    'date': date,
                    'duration': duration,
                    'work_entry_type_id': entry_type.id,
                    'l10n_ga_imported': True,
                }
                for date, duration in ii.spread_hours(hours, days)
            ]
        )

    def action_import(self):
        self.ensure_one()
        rows, issues = self._pipeline()
        slips = {slip.employee_id.id: slip for slip in self._draft_slips()}
        ctx = self._context()
        touched = self.env['hr.payslip']
        write_errors = []
        for row in rows:
            slip = slips[row.employee]
            try:
                with self.env.cr.savepoint():
                    for code, value in row.values:
                        kind, target = ctx.columns[code]
                        if kind == ii.AMOUNT:
                            self._write_amount(slip, self.env['hr.payslip.input.type'].browse(target), value)
                        else:
                            self._write_hours(slip, self.env['hr.work.entry.type'].browse(target), value)
            except ValueError:
                write_errors.append(
                    self.env._(
                        'Ligne %(row)s : %(employee)s, heures supérieures à la capacité de la période : ligne refusée.',
                        row=row.row,
                        employee=slip.employee_id.name,
                    )
                )
                continue
            touched |= slip
        if touched:
            touched._compute_worked_days_line_ids()
            touched.compute_sheet()
        summary = [self._issues_summary(issues), *write_errors]
        self.write(
            {
                'state': 'done',
                'valid_count': len(touched),
                'rejected_count': self._rejected_rows(issues) + len(write_errors),
                'issues_text': '\n'.join(text for text in summary if text) or False,
            }
        )
        return self._reopen()

    def _reopen(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
            'name': self.env._('Importer les variables du mois'),
        }
