import base64
import io
from datetime import date

import openpyxl

from odoo.exceptions import UserError
from odoo.tests import tagged

from .common import GaPayrollCase

SEPT = (date(2026, 9, 1), date(2026, 9, 30))


@tagged('post_install', '-at_install')
class TestInputImport(GaPayrollCase):
    """F3, patron 13, D-41 à D-43 : modèle du lot, aperçu, import des montants et des heures, recalcul."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        agreement = cls.env['l10n_ga.collective.agreement'].create(
            {
                'name': 'Convention import',
                'code': 'IMP',
                'company_id': cls.company.id,
                'overtime_rate_ids': [(0, 0, {'period': 'day', 'hours_from': 0, 'hours_to': 0, 'rate': 0.10})],
            }
        )
        values = {'l10n_ga_agreement_id': agreement.id, 'ssnid': 'CNSS', 'l10n_ga_nif': 'NIF'}
        cls.awa = cls._employee('Awa Ndong', 400_000, start=date(2020, 1, 1), registration_number='M001', **values)
        cls.paul = cls._employee('Paul Mba', 350_000, start=date(2020, 1, 1), registration_number='M002', **values)
        cls.outside = cls._employee('Hors lot', 300_000, start=date(2020, 1, 1), registration_number='M009')
        cls.pay_run = cls.env['hr.payslip.run'].create(
            {
                'name': 'Septembre 2026',
                'company_id': cls.company.id,
                'date_start': SEPT[0],
                'date_end': SEPT[1],
                'structure_id': cls.structure.id,
            }
        )
        for employee in (cls.awa, cls.paul):
            cls.env['hr.payslip'].create(
                {
                    'name': f'{employee.name} 09/2026',
                    'employee_id': employee.id,
                    'date_from': SEPT[0],
                    'date_to': SEPT[1],
                    'payslip_run_id': cls.pay_run.id,
                }
            )
        cls.pay_run.slip_ids.compute_sheet()

    def _slip(self, employee):
        return self.pay_run.slip_ids.filtered(lambda s: s.employee_id == employee)

    def _wizard(self, rows, header=None):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(header or ['Matricule', 'Salarié', 'GA_ASSID — Prime', 'GA_HS_J — Heures (heures)'])
        for row in rows:
            sheet.append(list(row))
        stream = io.BytesIO()
        workbook.save(stream)
        return self.env['l10n_ga.payslip.input.import'].create(
            {
                'payslip_run_id': self.pay_run.id,
                'import_file': base64.b64encode(stream.getvalue()),
                'import_filename': 'v.xlsx',
            }
        )

    def _imported(self, slip, code):
        return slip.input_line_ids.filtered(lambda line: line.code == code)

    def _hours(self, employee, code='GA_HS_J'):
        return self.env['hr.work.entry'].search(
            [('employee_id', '=', employee.id), ('code', '=', code), ('l10n_ga_imported', '=', True)]
        )

    # --- modèle ----------------------------------------------------------------------------------

    def test_template(self):
        wizard = self.env['l10n_ga.payslip.input.import'].create({'payslip_run_id': self.pay_run.id})
        wizard.action_generate_template()
        self.assertTrue(wizard.template_filename.endswith('.xlsx'))
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.template_file)))
        rows = list(workbook.active.iter_rows(values_only=True))
        header = rows[0]
        codes = [label.split(' — ')[0] for label in header[2:]]
        self.assertEqual(header[:2], ('Matricule', 'Salarié'))
        self.assertIn('GA_ASSID', codes)
        self.assertIn('GA_TRANSP', codes)
        self.assertNotIn('GA_LOAN', codes, 'les échéances de prêt ne se saisissent pas')
        self.assertEqual(
            [code for code in codes if code.startswith('GA_HS_')], ['GA_HS_J', 'GA_HS_N', 'GA_HS_DIM', 'GA_HS_FER']
        )
        self.assertEqual(len(codes), len(set(codes)))
        self.assertEqual(sorted(row[0] for row in rows[1:]), ['M001', 'M002'])

    def test_template_round_trip_imports_nothing(self):
        wizard = self.env['l10n_ga.payslip.input.import'].create({'payslip_run_id': self.pay_run.id})
        wizard.action_generate_template()
        wizard.import_file = wizard.template_file
        wizard.action_preview()
        self.assertEqual((wizard.valid_count, wizard.rejected_count, wizard.issues_text), (2, 0, False))

    # --- aperçu ----------------------------------------------------------------------------------

    def test_preview_rejects_wrong_rows_without_blocking_good_ones(self):
        wizard = self._wizard(
            [
                ('M001', 'Awa', 15_000, 7.5),
                ('M002', 'Paul', 'mille', None),
                ('M404', 'Inconnu', 1, None),
                ('M009', 'Hors lot', 1, None),
            ]
        )
        wizard.action_preview()
        self.assertEqual((wizard.state, wizard.valid_count, wizard.rejected_count), ('preview', 1, 3))
        self.assertIn('Ligne 3', wizard.issues_text)
        self.assertIn('n’est pas un nombre', wizard.issues_text)
        self.assertIn('M404', wizard.issues_text)
        self.assertIn('Hors lot', wizard.issues_text)
        self.assertFalse(self._imported(self._slip(self.awa), 'GA_ASSID'), 'l’aperçu n’écrit rien')

    def test_preview_duplicates_and_header_issues(self):
        wizard = self._wizard(
            [('M001', '', 1_000, None, 1), ('M001', '', 2_000, None, 2)],
            header=['Matricule', 'Salarié', 'GA_ASSID', 'GA_ASSID — bis', 'GA_INCONNU'],
        )
        wizard.action_preview()
        self.assertEqual((wizard.valid_count, wizard.rejected_count), (0, 2))
        self.assertIn('plusieurs lignes', wizard.issues_text)
        self.assertIn('GA_ASSID en double', wizard.issues_text)
        self.assertIn('GA_INCONNU inconnue', wizard.issues_text)

    def test_missing_key_column(self):
        wizard = self._wizard([('Awa', 1)], header=['Nom', 'GA_ASSID'])
        wizard.action_preview()
        self.assertEqual(wizard.valid_count, 0)
        self.assertIn('Matricule', wizard.issues_text)

    def test_unreadable_or_missing_file(self):
        wizard = self.env['l10n_ga.payslip.input.import'].create({'payslip_run_id': self.pay_run.id})
        with self.assertRaisesRegex(UserError, 'Choisissez'):
            wizard.action_preview()
        wizard.import_file = base64.b64encode(b'pas un classeur')
        with self.assertRaisesRegex(UserError, 'illisible'):
            wizard.action_preview()

    # --- import ----------------------------------------------------------------------------------

    def test_import_amounts_and_decimal_hours_then_recompute(self):
        wizard = self._wizard([('M001', 'Awa', '15 000', '7,5'), (None, 'Paul Mba', 20_000, None)])
        wizard.action_preview()
        wizard.action_import()
        self.assertEqual((wizard.state, wizard.valid_count, wizard.rejected_count), ('done', 2, 0))
        awa = self._slip(self.awa)
        prime = self._imported(awa, 'GA_ASSID')
        self.assertEqual((prime.amount, prime.l10n_ga_imported), (15_000, True))
        self.assertEqual(sum(self._hours(self.awa).mapped('duration')), 7.5)
        self.assertEqual(awa._l10n_ga_overtime_hours('day'), 7.5)
        totals = self._totals(awa)
        self.assertEqual(totals['GA_ASSID'], 15_000)
        self.assertTrue(totals['GA_HS_J'])
        self.assertEqual(self._totals(self._slip(self.paul))['GA_ASSID'], 20_000)

    def test_reimport_replaces_and_zero_deletes(self):
        slip = self._slip(self.awa)
        self.env['hr.payslip.input'].create(
            {'payslip_id': slip.id, 'input_type_id': self._input_type('GA_ASSID').id, 'amount': 999}
        )
        self._wizard([('M001', '', 15_000, 4)]).action_import()
        self._wizard([('M001', '', 10_000, 2.25)]).action_import()
        self.assertEqual(self._imported(slip, 'GA_ASSID').mapped('amount'), [10_000])
        self.assertEqual(self._hours(self.awa).mapped('duration'), [2.25])
        self._wizard([('M001', '', 0, 0)]).action_import()
        self.assertFalse(self._imported(slip, 'GA_ASSID'))
        self.assertFalse(self._hours(self.awa))
        self.assertNotIn('GA_ASSID', self._totals(slip))

    def test_empty_cell_leaves_value_unchanged(self):
        self._wizard([('M001', '', 15_000, None)]).action_import()
        self._wizard([('M001', '', None, 3)]).action_import()
        self.assertEqual(self._imported(self._slip(self.awa), 'GA_ASSID').amount, 15_000)

    def test_hours_spread_over_days_below_24h(self):
        self._wizard([('M002', '', None, 30.5)]).action_import()
        entries = self._hours(self.paul)
        self.assertEqual(sum(entries.mapped('duration')), 30.5)
        self.assertGreater(len(entries), 1)
        for day in set(entries.mapped('date')):
            day_total = sum(
                self.env['hr.work.entry']
                .search([('employee_id', '=', self.paul.id), ('date', '=', day)])
                .mapped('duration')
            )
            self.assertLessEqual(day_total, 24)
        self.assertEqual(max(entries.mapped('date')), SEPT[1], 'remplissage à partir du dernier jour')

    def test_hours_over_capacity_rejected_at_write(self):
        wizard = self._wizard([('M001', '', 5_000, 800), ('M002', '', 6_000, None)])
        wizard.action_import()
        self.assertEqual((wizard.valid_count, wizard.rejected_count), (1, 1))
        self.assertIn('capacité', wizard.issues_text)
        self.assertFalse(self._imported(self._slip(self.awa), 'GA_ASSID'), 'ligne refusée en entier')
        self.assertEqual(self._imported(self._slip(self.paul), 'GA_ASSID').amount, 6_000)
