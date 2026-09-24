import base64
import io
from collections import defaultdict
from datetime import date

import openpyxl

from odoo.exceptions import UserError, ValidationError
from odoo.tests import tagged

from ..report.xlsx_renderer import XlsxRenderer
from .common import GaPayrollCase

AUG = (date(2026, 8, 1), date(2026, 8, 31))
SEPT = (date(2026, 9, 1), date(2026, 9, 30))


@tagged('post_install', '-at_install')
class TestPayrollReports(GaPayrollCase):
    """F2, F13, D-49 à D-53 : livre de paie, état des charges, virements, billetage (valeurs, jamais de formules)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company.l10n_ga_cash_rounding = 500
        bgfi = cls.env['res.bank'].create({'name': 'BGFI'})
        uba = cls.env['res.bank'].create({'name': 'UBA'})
        cls.transfer_a = cls._with_bank(cls._employee('Virement A', 450_000, registration_number='V1'), bgfi, 'GA-1')
        cls.transfer_b = cls._with_bank(cls._employee('Virement B', 380_000, registration_number='V2'), uba, 'GA-2')
        cls.no_account = cls._employee('Sans compte', 300_000, registration_number='V3')
        cls.check = cls._employee('Chèque', 320_000, registration_number='C1')
        cls.check.version_id.l10n_ga_payment_mode = 'check'
        cls.cash = cls._employee('Espèces', 401_237, registration_number='E1')
        cls.cash.version_id.l10n_ga_payment_mode = 'cash'
        cls.cash_2 = cls._employee('Espèces 2', 277_777, registration_number='E2')
        cls.cash_2.version_id.l10n_ga_payment_mode = 'cash'
        cls.employees = cls.transfer_a | cls.transfer_b | cls.no_account | cls.check | cls.cash | cls.cash_2
        cls.pay_run = cls.env['hr.payslip.run'].create(
            {
                'name': 'Septembre',
                'company_id': cls.company.id,
                'date_start': SEPT[0],
                'date_end': SEPT[1],
                'structure_id': cls.structure.id,
            }
        )
        for employee in cls.employees:
            inputs = {'GA_TRANSP': 30_000} if employee == cls.transfer_a else None
            slip = cls._payslip(employee, *SEPT, inputs=inputs, compute_sheet=False)
            slip.payslip_run_id = cls.pay_run
        cls.pay_run.slip_ids.compute_sheet()
        cls.pay_run.slip_ids.action_payslip_done()
        cls.draft = cls._payslip(cls._employee('Brouillon', 999_999), *SEPT)

    @classmethod
    def _with_bank(cls, employee, bank, number):
        account = cls.env['res.partner.bank'].create(
            {'acc_number': number, 'partner_id': employee.work_contact_id.id, 'bank_id': bank.id}
        )
        employee.bank_account_ids = [(4, account.id)]
        return employee

    def _workbook(self, report_type, **values):
        wizard = self.env['l10n_ga.payroll.report'].create(
            {'company_id': self.company.id, 'report_type': report_type, 'payslip_run_id': self.pay_run.id, **values}
        )
        wizard.action_generate()
        self.assertTrue(wizard.filename.endswith('.xlsx'))
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.file)))

    @staticmethod
    def _rows(sheet):
        return list(sheet.iter_rows(values_only=True))

    def _net_pay(self, employees):
        slips = self.pay_run.slip_ids.filtered(lambda s: s.employee_id in employees)
        return sum(slips.line_ids.filtered(lambda line: line.code == 'GA_NET_PAY').mapped('total'))

    def assertNoFormula(self, workbook):
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertFalse(isinstance(cell.value, str) and cell.value.startswith('='), cell.coordinate)

    # --- livre de paie et état des charges -------------------------------------------------------

    def test_payroll_book_totals_equal_lines(self):
        workbook = self._workbook('book')
        self.assertNoFormula(workbook)
        rows = self._rows(workbook['Livre de paie'])
        header, body, total = rows[0], rows[1:-1], rows[-1]
        self.assertEqual(len(body), len(self.employees), 'brouillons exclus')
        self.assertNotIn('Brouillon', [row[1] for row in body])
        self.assertEqual(total[0], 'Total')
        codes = [label.split(' — ')[0] for label in header[4:]]
        expected = defaultdict(float)
        for line in self.pay_run.slip_ids.line_ids:
            expected[line.code] += line.total
        for index, code in enumerate(codes, start=4):
            self.assertAlmostEqual(total[index], sum(row[index] or 0 for row in body), delta=0.01, msg=code)
            self.assertAlmostEqual(total[index], expected[code], delta=0.01, msg=code)
        self.assertIn('NET', codes)
        self.assertIn('GA_TRANSP', codes)
        self.assertEqual({code for code, amount in expected.items() if amount} - set(codes), set())
        self.assertNotIn('GA_LOAN', codes, 'rubrique sans montant absente')

    def test_charges_statement(self):
        workbook = self._workbook('book')
        rows = self._rows(workbook['État des charges'])
        by_code = {row[0]: row for row in rows[1:] if row[0]}
        lines = self.pay_run.slip_ids.line_ids
        cnss_employee = -sum(lines.filtered(lambda line: line.code == 'GA_CNSS_SAL').mapped('total'))
        self.assertAlmostEqual(by_code['GA_CNSS_SAL'][2], cnss_employee, delta=0.01)
        self.assertEqual(by_code['GA_CNSS_SAL'][3], 0)
        pf = sum(lines.filtered(lambda line: line.code == 'GA_CNSS_PF').mapped('total'))
        self.assertAlmostEqual(by_code['GA_CNSS_PF'][3], pf, delta=0.01)
        for code in ('GA_CNSS_AT', 'GA_CNSS_AVID', 'GA_CNAMGS_PAT', 'GA_FNH', 'GA_CFP', 'GA_TCS', 'GA_IRPP'):
            self.assertIn(code, by_code)
        self.assertNotIn('NET', by_code)
        total = rows[-1]
        details = [row for row in rows[1:-1] if row[0]]
        self.assertAlmostEqual(total[2], sum(row[2] for row in details), delta=0.01)
        self.assertAlmostEqual(total[3], sum(row[3] for row in details), delta=0.01)
        self.assertAlmostEqual(total[4], total[2] + total[3], delta=0.01)

    # --- virements -------------------------------------------------------------------------------

    def test_transfers_one_sheet_per_bank(self):
        workbook = self._workbook('transfer')
        self.assertNoFormula(workbook)
        self.assertEqual(sorted(workbook.sheetnames), ['BGFI', 'Chèques', 'Sans compte', 'UBA'])
        bgfi = self._rows(workbook['BGFI'])
        self.assertEqual([row[:4] for row in bgfi[1:-1]], [('V1', 'Virement A', 'BGFI', 'GA-1')])
        self.assertEqual(bgfi[-1][4], self._net_pay(self.transfer_a))
        self.assertEqual(self._rows(workbook['Chèques'])[-1][4], self._net_pay(self.check))
        self.assertEqual(self._rows(workbook['Sans compte'])[1][0], 'V3')
        paid = sum(self._rows(workbook[name])[-1][4] for name in workbook.sheetnames)
        self.assertEqual(paid, self._net_pay(self.employees - self.cash - self.cash_2))

    def test_transfers_use_frozen_account(self):
        self.transfer_a.bank_account_ids.acc_number = 'GA-CHANGE'
        rows = self._rows(self._workbook('transfer')['BGFI'])
        self.assertEqual(rows[1][3], 'GA-1')

    # --- billetage -------------------------------------------------------------------------------

    def test_cash_breakdown_sum_equals_cash_net_pay(self):
        workbook = self._workbook('cash')
        self.assertNoFormula(workbook)
        rows = self._rows(workbook['Billetage'])
        header = rows[0]
        denominations = [int(value) for value in header[3:-1]]
        self.assertEqual(denominations, [10_000, 5_000, 2_000, 1_000, 500])
        body, counts, amounts = rows[1:-2], rows[-2], rows[-1]
        self.assertEqual([row[0] for row in body], ['E1', 'E2'])
        for row in body:
            self.assertEqual(sum(n * d for n, d in zip(row[3:-1], denominations, strict=True)) + row[-1], row[2])
        cash_pay = self._net_pay(self.cash | self.cash_2)
        self.assertEqual(sum(amounts[3:-1]) + amounts[-1], cash_pay)
        self.assertEqual(counts[2], cash_pay)
        self.assertEqual(amounts[-1], 0, 'paies arrondies à 500 : aucune pièce')
        self.assertEqual(list(counts[3:-1]), [sum(row[i] for row in body) for i in range(3, len(header) - 1)])

    # --- sélection -------------------------------------------------------------------------------

    def test_selection_by_period_and_company(self):
        august = self._validated(self.transfer_a, *AUG)
        workbook = self._workbook('book', payslip_run_id=False, date_from=AUG[0], date_to=AUG[1])
        body = self._rows(workbook['Livre de paie'])[1:-1]
        self.assertEqual([row[1] for row in body], ['Virement A'])
        self.assertEqual(body[0][2], '08/2026')
        self.assertTrue(august.l10n_ga_frozen_date)
        other = self.env['res.company'].create({'name': 'Autre'})
        wizard = self.env['l10n_ga.payroll.report'].create(
            {'company_id': other.id, 'date_from': SEPT[0], 'date_to': SEPT[1]}
        )
        with self.assertRaisesRegex(UserError, 'Aucun bulletin'):
            wizard.action_generate()

    def test_period_required(self):
        with self.assertRaises(ValidationError):
            self.env['l10n_ga.payroll.report'].create({'company_id': self.company.id, 'date_from': False})
        with self.assertRaises(ValidationError):
            self.env['l10n_ga.payroll.report'].create(
                {'company_id': self.company.id, 'date_from': SEPT[1], 'date_to': SEPT[0]}
            )

    def test_run_button(self):
        action = self.pay_run.action_l10n_ga_reports()
        self.assertEqual(action['context']['default_payslip_run_id'], self.pay_run.id)

    # --- constructeur ----------------------------------------------------------------------------

    def test_renderer_sheet_names_and_values(self):
        renderer = XlsxRenderer()
        renderer.add_sheet('Banque [test]: ?/*', ['A', 'B'], [['x', 1.5], [None, True]], totals=[['T', 1.5]])
        renderer.add_sheet('Banque [test]: ?/*', ['A'], [])
        renderer.add_sheet('N' * 40, ['A'], [])
        workbook = openpyxl.load_workbook(io.BytesIO(renderer.build()))
        self.assertEqual(workbook.sheetnames, ['Banque  test', 'Banque  test (2)', 'N' * 31])
        rows = self._rows(workbook.worksheets[0])
        self.assertEqual(rows, [('A', 'B'), ('x', 1.5), (None, None), ('T', 1.5)])
